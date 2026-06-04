from pypdf import PdfReader

from docx import Document as DocxDocument

from pptx import Presentation

from llama_index.core import (
    Document,
    VectorStoreIndex,
    StorageContext
)

from llama_index.core.node_parser import (
    SentenceSplitter
)

from llama_index.embeddings.openai import (
    OpenAIEmbedding
)

from llama_index.vector_stores.postgres import (
    PGVectorStore
)

import os


class IngestionService:

    def clean_text(
        self,
        text: str
    ) -> str:

        if not text:
            return ""

        return (
            text
            .replace("\x00", "")
            .replace("\u0000", "")
            .strip()
        )

    def extract_pdf_text(
        self,
        file_path: str
    ) -> str:

        reader = PdfReader(
            file_path
        )

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text += (
                    page_text
                    + "\n"
                )

        return self.clean_text(
            text
        )

    def extract_docx_text(
        self,
        file_path: str
    ) -> str:

        document = DocxDocument(
            file_path
        )

        text = ""

        for paragraph in document.paragraphs:

            text += (
                paragraph.text
                + "\n"
            )

        return self.clean_text(
            text
        )

    def extract_txt_text(
        self,
        file_path: str
    ) -> str:

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as f:

            text = f.read()

        return self.clean_text(
            text
        )

    def extract_pptx_text(
        self,
        file_path: str
    ) -> str:

        presentation = Presentation(
            file_path
        )

        text = ""

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(
                    shape,
                    "text"
                ):

                    text += (
                        shape.text
                        + "\n"
                    )

        return self.clean_text(
            text
        )

    def extract_text(
        self,
        file_path: str
    ) -> str:

        extension = os.path.splitext(
            file_path
        )[1].lower()

        if extension == ".pdf":

            return self.extract_pdf_text(
                file_path
            )

        elif extension == ".docx":

            return self.extract_docx_text(
                file_path
            )

        elif extension == ".txt":

            return self.extract_txt_text(
                file_path
            )

        elif extension == ".pptx":

            return self.extract_pptx_text(
                file_path
            )

        else:

            raise Exception(
                f"Unsupported file type: {extension}"
            )

    def ingest_file(
        self,
        file_path: str
    ):

        print(
            f"Processing file: {file_path}"
        )

        # -------------------------
        # Extract Text
        # -------------------------

        text = self.extract_text(
            file_path
        )

        if not text:

            raise Exception(
                "No text found in document"
            )

        print(
            "First 500 characters:"
        )

        print(
            text[:500]
        )

        # -------------------------
        # Create LlamaIndex Document
        # -------------------------

        extension = os.path.splitext(
            file_path
        )[1].lower()

        documents = [
            Document(
                text=text,
                metadata={
                    "file_name":
                        os.path.basename(
                            file_path
                        ),
                    "file_type":
                        extension
                }
            )
        ]

        # -------------------------
        # Chunking
        # -------------------------

        parser = SentenceSplitter(
            chunk_size=512,
            chunk_overlap=50
        )

        nodes = parser.get_nodes_from_documents(
            documents
        )

        print(
            f"Created {len(nodes)} chunks"
        )

        for node in nodes:

            node.text = self.clean_text(
                node.text
            )

        # -------------------------
        # Embedding Model
        # -------------------------

        embed_model = OpenAIEmbedding(
            model="text-embedding-3-small"
        )

        # -------------------------
        # PGVector
        # -------------------------

        vector_store = PGVectorStore.from_params(
            host=os.getenv(
                "DB_HOST"
            ),
            port=5432,
            database=os.getenv(
                "DB_NAME"
            ),
            user=os.getenv(
                "DB_USER"
            ),
            password=os.getenv(
                "DB_PASSWORD"
            ),
            table_name="resume_chunks",
            embed_dim=1536
        )

        storage_context = (
            StorageContext.from_defaults(
                vector_store=vector_store
            )
        )

        # -------------------------
        # Save Embeddings
        # -------------------------

        VectorStoreIndex(
            nodes=nodes,
            storage_context=storage_context,
            embed_model=embed_model
        )

        print(
            "Embeddings saved successfully"
        )

        return len(nodes)